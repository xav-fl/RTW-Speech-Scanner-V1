"""Ingest documents you supply yourself, instead of scraping.

This is the recommended path when you have already chosen your documents, or
when a source cannot be reached programmatically. It produces exactly the same
`Document` objects the harvesters do, so everything downstream is identical.

Two steps:

    eps ingest --dir docs/ --write-template manifest.csv   # scaffold
    # ...fill in institution, url, title, doc_type, date...
    eps ingest --run-dir runs/x --manifest manifest.csv    # build

The manifest is the point: metadata that a scraper would read from page markup
has to come from somewhere, and having you state it explicitly is more reliable
than having the program guess. The manifest is also a citable record of exactly
which documents entered the dataset.
"""

from __future__ import annotations

import csv
import re
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Iterable, Optional

from ..models import DELIVERY_STATUS, DOC_TYPES, Document, sha256
from ..normalise import (
    SEGMENTER_VERSION, html_to_text, make_doc_id, mark_attribution,
    normalise_text, segment,
)
from .store import RawStore, canonical_url

MANIFEST_COLUMNS = [
    "path", "institution", "country", "url", "title", "doc_type", "date",
    "delivery_status", "speaker_hint", "notes",
]

TEXT_SUFFIXES = {".txt", ".md", ".text"}
HTML_SUFFIXES = {".html", ".htm", ".xhtml"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}

SUPPORTED = TEXT_SUFFIXES | HTML_SUFFIXES | PDF_SUFFIXES | DOCX_SUFFIXES

_AS_PREPARED = re.compile(r"as\s+prepared\s+for\s+delivery|prepared\s+remarks", re.I)
_AS_DELIVERED = re.compile(r"as\s+delivered|\btranscript\b", re.I)


class IngestError(Exception):
    pass


# --- text extraction -------------------------------------------------------


def extract_text(path: Path) -> str:
    """Extract paragraph-preserving text from a local file."""
    suffix = path.suffix.lower()
    raw = path.read_bytes()

    if suffix in HTML_SUFFIXES:
        return html_to_text(raw.decode("utf-8", errors="replace"))

    if suffix in TEXT_SUFFIXES:
        return normalise_text(raw.decode("utf-8", errors="replace"))

    if suffix in PDF_SUFFIXES:
        return _pdf_to_text(path)

    if suffix in DOCX_SUFFIXES:
        return _docx_to_text(path)

    raise IngestError(
        f"unsupported file type {suffix!r}; supported: {sorted(SUPPORTED)}"
    )


def _pdf_to_text(path: Path) -> str:
    """PDFs need care: a naive extract runs lines together and destroys the
    paragraph boundaries the unit scheme depends on. Blank-line gaps and
    line-ending punctuation are used to rebuild paragraphs."""
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise IngestError("pdfplumber is required to ingest PDFs") from exc

    pages: list[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text() or "")
    raw = "\n".join(pages)

    paragraphs: list[list[str]] = [[]]
    for line in raw.split("\n"):
        stripped = line.strip()
        if not stripped:
            if paragraphs[-1]:
                paragraphs.append([])
            continue
        # A short line ending in sentence punctuation usually closes a block.
        paragraphs[-1].append(stripped)
        if re.search(r"[.!?][\"'\u2019\u201d]?$", stripped) and len(stripped) < 60:
            paragraphs.append([])

    blocks = [" ".join(p) for p in paragraphs if p]
    return normalise_text("\n\n".join(blocks))


def _docx_to_text(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise IngestError("python-docx is required to ingest .docx files") from exc
    document = docx.Document(str(path))
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return normalise_text("\n\n".join(blocks))


# --- manifest --------------------------------------------------------------


def write_template(paths: Iterable[Path], out: Path, country: str = "") -> int:
    """Scaffold a manifest for a folder of documents, for you to fill in."""
    out.parent.mkdir(parents=True, exist_ok=True)
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=MANIFEST_COLUMNS)
        w.writeheader()
        for p in sorted(paths):
            if p.suffix.lower() not in SUPPORTED:
                continue
            w.writerow(
                {
                    "path": str(p),
                    "institution": "",
                    "country": country,
                    "url": "",
                    "title": p.stem.replace("_", " ").replace("-", " "),
                    "doc_type": "speech",
                    "date": "",
                    "delivery_status": "unknown",
                    "speaker_hint": "",
                    "notes": "",
                }
            )
            n += 1
    return n


def read_manifest(path: Path) -> list[dict]:
    # utf-8-sig: Excel on Windows often saves CSV with a byte-order mark at the
    # start of the file. Reading as plain utf-8 leaves that mark glued onto the
    # first header cell, so "path" is silently read as something else and the
    # file appears to be "missing" a column that is, in fact, right there.
    # utf-8-sig strips the mark if present and behaves exactly like utf-8 if not.
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    if not rows:
        raise IngestError(f"{path} is empty")
    missing = {"path", "institution", "country", "url", "date"} - set(rows[0])
    if missing:
        raise IngestError(f"{path} is missing required columns: {sorted(missing)}")
    return rows


def validate_row(row: dict, line: int) -> list[str]:
    """Fail loudly on bad metadata rather than silently producing bad rows."""
    problems: list[str] = []
    p = Path(row.get("path", "").strip())
    if not row.get("path", "").strip():
        problems.append("path is empty")
    elif not p.exists():
        problems.append(f"file not found: {p}")
    elif p.suffix.lower() not in SUPPORTED:
        problems.append(f"unsupported type {p.suffix!r}")

    if not row.get("institution", "").strip():
        problems.append("institution is empty (needed for the officeholder join)")
    if row.get("country", "").strip() not in {"US", "UK"}:
        problems.append("country must be US or UK")
    if not row.get("url", "").strip():
        problems.append("url is empty (needed for the dataset's url column)")

    d = row.get("date", "").strip()
    if not d:
        problems.append("date is empty (needed for speaker eligibility)")
    else:
        try:
            date.fromisoformat(d)
        except ValueError:
            problems.append(f"date {d!r} is not ISO format (YYYY-MM-DD)")

    dt = (row.get("doc_type") or "speech").strip()
    if dt not in DOC_TYPES:
        problems.append(f"doc_type {dt!r} not in {list(DOC_TYPES)}")

    ds = (row.get("delivery_status") or "unknown").strip()
    if ds not in DELIVERY_STATUS:
        problems.append(f"delivery_status {ds!r} not in {list(DELIVERY_STATUS)}")

    return [f"row {line}: {m}" for m in problems]


# --- ingestion -------------------------------------------------------------


def ingest_row(row: dict, store: Optional[RawStore] = None) -> Document:
    path = Path(row["path"].strip())
    raw_bytes = path.read_bytes()
    text = extract_text(path)
    if not text.strip():
        raise IngestError(f"{path}: no text extracted")

    doc_type = (row.get("doc_type") or "speech").strip()
    url = canonical_url(row["url"].strip())

    if store is not None:
        digest = store.put(
            url, raw_bytes,
            {"institution": row["institution"].strip(), "source_index": "local_ingest",
             "local_path": str(path)},
        )
    else:
        digest = sha256(raw_bytes)

    doc_id = make_doc_id(url, digest)
    units = mark_attribution(segment(text, doc_id), doc_type)

    delivery = (row.get("delivery_status") or "").strip() or "unknown"
    if delivery == "unknown":
        haystack = f"{row.get('title','')} {text[:2000]}"
        if _AS_PREPARED.search(haystack):
            delivery = "as_prepared"
        elif _AS_DELIVERED.search(haystack):
            delivery = "as_delivered"

    return Document(
        doc_id=doc_id,
        institution=row["institution"].strip(),
        country=row["country"].strip(),
        url=url,
        title=(row.get("title") or "").strip(),
        doc_type=doc_type,
        pub_date=date.fromisoformat(row["date"].strip()),
        delivery_status=delivery,
        speaker_hint=(row.get("speaker_hint") or "").strip() or None,
        raw_sha256=digest,
        text_sha256=sha256(text),
        segmenter_version=SEGMENTER_VERSION,
        retrieved_at=datetime.now(timezone.utc),
        source_index="local_ingest",
        text=text,
        units=units,
    )


def ingest_manifest(
    manifest_path: Path, store: Optional[RawStore] = None, strict: bool = True
) -> tuple[list[Document], list[str]]:
    rows = read_manifest(manifest_path)

    problems: list[str] = []
    for i, row in enumerate(rows, start=2):  # line 1 is the header
        problems.extend(validate_row(row, i))
    if problems and strict:
        raise IngestError(
            "manifest validation failed:\n  " + "\n  ".join(problems)
        )

    docs: list[Document] = []
    errors: list[str] = list(problems)
    for i, row in enumerate(rows, start=2):
        if any(f"row {i}:" in p for p in problems):
            continue
        try:
            docs.append(ingest_row(row, store))
        except Exception as exc:
            errors.append(f"row {i}: {exc}")
    return docs, errors
